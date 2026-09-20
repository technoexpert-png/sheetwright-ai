#!/usr/bin/env python3
"""Regenerate the sample spreadsheets used to exercise the ingestion pipeline.

Run from anywhere:

    .venv/bin/python samples/make_samples.py            # write both files
    .venv/bin/python samples/make_samples.py --verify    # write, then dump the
                                                         # raw grids back out

The script is deterministic and idempotent: every value below is a literal, no
timestamps or random data are used, so re-running it produces byte-identical
files and a clean `git diff`.

The two samples deliberately stress *different* failure modes, so that a fix
for one does not accidentally paper over the other:

  * sample_a_contacts.csv -- SEMANTIC mess. The grid is rectangular and
    well-formed; what is hard is deciding what the columns *mean* and cleaning
    the values inside them.
  * sample_b_roster.xlsx  -- STRUCTURAL mess. The column names are closer to
    the canonical schema; what is hard is recovering a flat table at all
    (merged bands, spacer rows, duplicate headers, Excel type coercion).

Canonical target schema (see app/models.py::CANONICAL_FIELDS):
    full_name (required), email (required), company (optional), phone (optional)
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import re
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

# Resolve paths relative to this file, not the caller's cwd, so the script can
# be invoked from the repo root or from inside samples/ with the same result.
SAMPLES_DIR = Path(__file__).resolve().parent
CSV_PATH = SAMPLES_DIR / "sample_a_contacts.csv"
XLSX_PATH = SAMPLES_DIR / "sample_b_roster.xlsx"


# ---------------------------------------------------------------------------
# FILE A -- sample_a_contacts.csv
# ---------------------------------------------------------------------------
# Failure modes exercised: junk preamble above the header, non-standard but
# recognizable header names, one unmappable column, one genuinely ambiguous
# column, and dirty values (missing/malformed email, four phone formats, a
# fully blank line mid-file, untrimmed whitespace, "LAST, First" name order).

# Two junk lines the header detector has to skip before row 3.
CSV_PREAMBLE: list[list[str]] = [
    ["Q3 Outreach List — CONFIDENTIAL"],
    ["Exported 2026-09-14"],
]

# None of these match the canonical names exactly. `Notes` is unmappable and
# `Secondary Contact` is ambiguous on purpose -- its values are a mix of names
# and phone numbers, so content sniffing cannot settle it either.
CSV_HEADER: list[str] = [
    "Contact",
    "E-mail Addr",
    "Org",
    "Cell #",
    "Notes",
    "Secondary Contact",
]

# An empty list writes a genuinely blank line (just the line terminator), which
# csv.reader hands back as `[]` -- a harsher case than a row of empty strings.
BLANK_ROW: list[str] = []

CSV_ROWS: list[list[str]] = [
    # 1. Clean baseline. Phone format 1 of 4: (555) 010-1234.
    #    `Secondary Contact` holds a person's name here...
    ["Ana Ruiz", "ana.ruiz@northwind.example", "Northwind Traders",
     "(555) 010-1234", "Met at the expo", "Diego Ruiz"],
    # 2. ...and a phone number here. That inconsistency is what makes the
    #    column ambiguous rather than merely unusual. Phone format 2: dots.
    ["Brian Okafor", "b.okafor@contoso.example", "Contoso Ltd",
     "555.010.1234", "", "555.010.9876"],
    # 3. Required field missing: no email at all.  Phone format 3: E.164-ish.
    ["Chen Wei", "", "Fabrikam",
     "+1 555 010 1234", "No email on the business card", ""],
    # 4. Malformed email (double @, no TLD) -- present but invalid, which is a
    #    different diagnostic from row 3's absent value. Phone format 4: bare
    #    ten digits.
    ["Jane Doe", "jane.doe@@example", "Adventure Works",
     "5550101234", "typo carried over from the CRM", "Ops desk"],
    # 5. Leading/trailing whitespace in several cells, plus a doubled inner
    #    space in the name and a whitespace-only `Secondary Contact`.
    ["  Priya   Nair  ", "  priya.nair@tailspin.example ", " Tailspin Toys",
     " (555) 010-7788 ", " follow up in October ", " "],
    # 6. Completely blank row in the middle of the data -- must not terminate
    #    parsing and must not emit an all-empty record.
    BLANK_ROW,
    # 7. "LASTNAME, Firstname" order; the surname is also shouted.
    ["NAKAMURA, Kenji", "k.nakamura@litware.example", "Litware Inc",
     "+1 555 010 4455", "Prefers email", "NAKAMURA, Yuki"],
    # 8. Both optional fields empty; required fields present. Should still be a
    #    complete row.
    ["Sofia Almeida", "sofia@wingtip.example", "",
     "", "Company unknown", ""],
    # 9. Phone with an extension -- a fifth shape that must survive
    #    normalization without losing the "ext. 14" part silently.
    ["Marcus Webb", "marcus.webb@proseware.example", "Proseware",
     "555-010-2200 ext. 14", "Call after 3pm", "+1 555 010 2201"],
    # 10. Honorific + suffix around the name, and a company containing a comma
    #     (so the CSV writer must quote it and the reader must not split it).
    ["Dr. Aiko Tanaka, PhD", "a.tanaka@vanarsdel.example", "Van Arsdel, Ltd.",
     "5550103399", "", "Aiko Tanaka"],
]


def write_csv(path: Path = CSV_PATH) -> Path:
    """Write FILE A. Overwrites in place; no state is read back in."""
    # newline="" is required by the csv module so it controls line endings
    # itself (it emits \r\n, which is what real spreadsheet exports look like).
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerows(CSV_PREAMBLE)
        writer.writerow(CSV_HEADER)
        writer.writerows(CSV_ROWS)
    return path


# ---------------------------------------------------------------------------
# FILE B -- sample_b_roster.xlsx
# ---------------------------------------------------------------------------
# Failure modes exercised: a merged title band, a second merged range *inside*
# the data whose value must be propagated down, header casing/spacing noise, a
# duplicate `Email` column, a blank spacer row, and Excel-native typing
# (numeric phones, real date cells).
#
# Sheet layout (1-based):
#   row 1      merged A1:G1 title band
#   row 2      header row
#   row 3      blank spacer
#   rows 4-11  eight data rows; A4:A6 is a vertically merged Region cell

XLSX_TITLE = "Partner Roster — Ingest Test"

# `  full name ` keeps its leading/trailing spaces, `EMAIL` is shouted, and the
# second `Email` at the end is the duplicate.
XLSX_HEADER: list[str] = [
    "Region",
    "  full name ",
    "EMAIL",
    "Employer",
    "Telephone",
    "Joined",
    "Email",
]

# Values only; the merged Region cell is written once (top-left) and left blank
# on the two rows beneath, which is exactly how Excel stores a vertical merge.
XLSX_ROWS: list[list[object]] = [
    # row 4 -- top-left of the merged A4:A6 Region band. Real date cell.
    ["EMEA", "Lena Fischer", "lena.fischer@northwind.example", "Northwind Traders",
     "+44 20 7946 0958", dt.date(2024, 3, 11), "lena.fischer@northwind.example"],
    # row 5 -- Region is blank because of the merge; the value "EMEA" must be
    #          propagated here. Phone is a NUMBER with a fractional part (an
    #          extension that Excel swallowed into the value), so openpyxl
    #          hands it back as the float 5550102200.14 -- naive str() then
    #          yields "5550102200.14" instead of a phone number.
    [None, "Tomas Novak", "t.novak@contoso.example", "Contoso Ltd",
     5550102200.14, dt.date(2023, 11, 2), "t.novak@contoso.example"],
    # row 6 -- last row of the merged Region band.
    [None, "Amara Singh", "amara.singh@fabrikam.example", "Fabrikam",
     "(555) 010-8822", dt.date(2025, 6, 30), "amara.singh@fabrikam.example"],
    # row 7 -- company only; both required fields are missing.
    ["APAC", None, None, "Yamato Logistics", None, None, None],
    # row 8 -- the two email columns DISAGREE (different domain), so the
    #          duplicate-header resolution is observable in the output.
    ["APAC", "Hiroshi Sato", "h.sato@litware.example", "Litware Inc",
     "555.010.3131", dt.date(2022, 9, 19), "hiroshi.sato@litware.co.example"],
    # row 9 -- phone as a whole NUMBER. openpyxl returns int here (it writes
    #          integral numerics without a decimal point), and the cell also
    #          carries a scientific number format, so Excel *displays* it as
    #          1.56E+10 while the stored value is 15550109090. The leading
    #          country code has already lost its "+". The join date is a
    #          *string* in ambiguous dd/mm-vs-mm/dd form, sitting next to the
    #          real date cells above.
    ["AMER", "Grace Adeyemi", "grace@wingtip.example", "Wingtip Toys",
     15550109090, "07/01/2025", "grace@wingtip.example"],
    # row 10 -- "LASTNAME, Firstname", an extension in the phone, and a
    #           datetime carrying a time component.
    ["AMER", "ORTIZ, Manuel", "m.ortiz@proseware.example", "Proseware",
     "+1 (555) 010-5566 x21", dt.datetime(2024, 12, 2, 9, 30), "m.ortiz@proseware.example"],
    # row 11 -- untrimmed name, shouted email, empty phone, and an empty
    #           duplicate-email cell while the primary is populated.
    ["EMEA", "  Yusuf   Demir ", "YUSUF.DEMIR@TAILSPIN.EXAMPLE", "Tailspin Toys",
     None, dt.date(2021, 4, 5), None],
]

# Frozen document timestamp -- see write_xlsx() and _freeze_xlsx().
FIXED_TIMESTAMP = dt.datetime(2026, 1, 1, 0, 0, 0)
FIXED_ZIP_TIME = FIXED_TIMESTAMP.timetuple()[:6]

TITLE_MERGE = "A1:G1"    # the top band
REGION_MERGE = "A4:A6"   # the merge inside the data region
HEADER_ROW = 2
FIRST_DATA_ROW = 4       # row 3 is the deliberate blank spacer


def write_xlsx(path: Path = XLSX_PATH) -> Path:
    """Write FILE B. Rebuilt from scratch each run, so it is idempotent."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Roster"

    # Pin the document properties. openpyxl otherwise stamps "created" and
    # "modified" with the current time, which would make every regeneration a
    # spurious binary diff in git.
    wb.properties.creator = "technoexpert-ingest samples"
    wb.properties.lastModifiedBy = "technoexpert-ingest samples"
    wb.properties.created = FIXED_TIMESTAMP
    wb.properties.modified = FIXED_TIMESTAMP

    # Title band: write the value into the top-left cell, then merge. openpyxl
    # discards anything written to the other cells of a merged range, so order
    # matters here.
    ws["A1"] = XLSX_TITLE
    ws.merge_cells(TITLE_MERGE)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    # Header row (row 2). Row 3 is left entirely untouched -> blank spacer.
    for col, name in enumerate(XLSX_HEADER, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=name)
        cell.font = Font(bold=True)

    # Data rows.
    for offset, row_values in enumerate(XLSX_ROWS):
        row_idx = FIRST_DATA_ROW + offset
        for col, value in enumerate(row_values, start=1):
            if value is None:
                continue  # leave truly empty, don't write an empty string
            ws.cell(row=row_idx, column=col, value=value)

    # Row 9's numeric phone gets Excel's scientific format, so a human opening
    # the file sees 1.56E+10 even though the stored value is exact. The parser
    # must read the value, not the rendering.
    ws["E9"].number_format = "0.00E+00"

    # The vertical Region merge inside the data. A4 already holds "EMEA"; A5
    # and A6 were skipped above, which is what the merge expects.
    ws.merge_cells(REGION_MERGE)
    ws["A4"].alignment = Alignment(vertical="top")

    # Cosmetic only -- makes the file readable if a human opens it.
    for col, width in zip("ABCDEFG", (10, 22, 34, 22, 24, 14, 34)):
        ws.column_dimensions[col].width = width

    wb.save(path)
    _freeze_xlsx(path)
    return path


def _freeze_xlsx(path: Path) -> None:
    """Rewrite the saved workbook so two runs produce byte-identical files.

    An .xlsx is a zip, and two things inside it would otherwise change on every
    run: the zip entries' modification times, and `dcterms:modified` in
    docProps/core.xml, which openpyxl stamps with `now` at save time regardless
    of what `wb.properties.modified` was set to. Both are pinned here so the
    committed sample does not churn in git.
    """
    stamp = FIXED_TIMESTAMP.strftime("%Y-%m-%dT%H:%M:%SZ")
    with zipfile.ZipFile(path) as src:
        entries = [(info, src.read(info.filename)) for info in src.infolist()]

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for info, data in entries:
            if info.filename == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>" + stamp.encode() + rb"\g<2>",
                    data,
                )
            # A fresh ZipInfo drops the original timestamp; carry over only the
            # fields that actually describe the entry.
            frozen = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_TIME)
            frozen.compress_type = info.compress_type
            frozen.external_attr = info.external_attr
            out.writestr(frozen, data)


# ---------------------------------------------------------------------------
# Verification -- read the files back and dump the raw grids.
# ---------------------------------------------------------------------------
# This exists so the structural claims above (merged ranges, blank rows,
# duplicate headers, numeric phones) can be confirmed from the files on disk
# rather than trusted from the code that wrote them.


def dump_csv(path: Path = CSV_PATH) -> None:
    print(f"\n=== {path.name} (csv.reader raw rows) ===")
    with path.open("r", encoding="utf-8", newline="") as fh:
        for i, row in enumerate(csv.reader(fh), start=1):
            marker = "  <-- BLANK LINE" if not row else ""
            print(f"  row {i:>2}: {row!r}{marker}")


def dump_xlsx(path: Path = XLSX_PATH) -> None:
    print(f"\n=== {path.name} (openpyxl raw cells) ===")
    wb = load_workbook(path)
    ws = wb.active
    print(f"  sheet={ws.title!r} dims={ws.dimensions} "
          f"max_row={ws.max_row} max_col={ws.max_column}")
    print(f"  merged ranges: {[str(r) for r in ws.merged_cells.ranges]}")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        idx = row[0].row
        cells = [f"{c.value!r}:{type(c.value).__name__}" for c in row]
        blank = all(c.value is None for c in row)
        marker = "  <-- BLANK ROW" if blank else ""
        print(f"  row {idx:>2}: [{', '.join(cells)}]{marker}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--verify",
        action="store_true",
        help="after writing, load both files back and print the raw grids",
    )
    args = parser.parse_args()

    for path in (write_csv(), write_xlsx()):
        print(f"wrote {path}")

    if args.verify:
        dump_csv()
        dump_xlsx()


if __name__ == "__main__":
    main()
